#!/usr/bin/env python3
"""第三轮高难度回归：单页桥 actions、xlsx 导出、webhook 通知、jobs 编排、
表格解析、LLM 打桩、嵌套 jpath、deny+same_domain、配置校验拦截、fetch 组合。
用法: python3 tests/run_tests3.py
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))

from mock_sites import Handler  # noqa: E402
from universal_scraper.engine_v3 import run_task  # noqa: E402
from universal_scraper.config import ConfigError, validate_task  # noqa: E402

PY = sys.executable
TMP = ROOT / "outputs" / ".test_tmp"
PASS, FAIL = [], []


def make_task(name: str, cfg: dict) -> Path:
    d = TMP / name
    (d / "modules").mkdir(parents=True, exist_ok=True)
    (d / "config.json").write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    return d


def cleanup(name: str):
    import glob, shutil
    for pat in (f"outputs/{name}.*", f"outputs/items/{name}.*",
                f"outputs/.state_{name}.json", f"outputs/.pending_{name}.json",
                f"outputs/.run_{name}.log", f"outputs/.seen_{name}.txt"):
        for f in glob.glob(str(ROOT / pat)):
            try:
                os.remove(f)
            except FileNotFoundError:
                pass
    shutil.rmtree(TMP / name, ignore_errors=True)


def read_items(name: str) -> list:
    fp = ROOT / "outputs" / "items" / f"{name}.jsonl"
    if not fp.exists():
        return []
    return [json.loads(x) for x in fp.read_text(encoding="utf-8").splitlines() if x.strip()]


def check(name: str, cond: bool, detail: str = ""):
    if cond:
        PASS.append(name)
        print(f"  ✅ {name}" + (f"  {detail}" if detail else ""))
    else:
        FAIL.append(name)
        print(f"  ❌ {name}  {detail}")


def main():
    srv = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    print(f"mock server: {base}\n")

    # 1. 单页桥 actions（pool=false）
    print("[1/10] 浏览器单页桥（pool=false）+ 动作链")
    name = "v1_single_actions"
    cfg = {
        "name": name, "start_urls": [f"{base}/actions.html"],
        "queue": {"max_depth": 1, "max_requests": 5, "max_concurrency": 1},
        "source": {"type": "browser", "pool": False, "remove_overlays": True,
                   "actions": [{"type": "click", "selector": "#more", "ms": 400}]},
        "rules": [{"match": "contains", "pattern": "/actions", "parser": "js"}],
        "parsers": {"js": {"type": "html", "row_css": ".item", "fields": {"title": {"css": "span.t"}}}},
        "pipelines": [{"type": "filter", "field": "title", "op": "non_empty"}],
        "storage": {"type": "jsonl", "name": name},
        "output": {"dir": "outputs", "base_name": name},
        "anti_bot": {"min_interval": 0.02, "max_retries": 1},
    }
    d = make_task(name, cfg)
    try:
        r = run_task(d)
        items = read_items(name)
        check(name, r["total"] == 6 and any(i["title"] == "loaded-1" for i in items),
              f"total={r['total']}")
    finally:
        cleanup(name)

    # 2. xlsx 导出完整性
    print("[2/10] xlsx 导出完整性（openpyxl 校验行数）")
    name = "v2_xlsx"
    cfg = {
        "name": name, "start_urls": [f"{base}/p1.html", f"{base}/p2.html"],
        "queue": {"max_depth": 1, "max_requests": 10, "max_concurrency": 2},
        "source": {"type": "http"},
        "rules": [{"match": "contains", "pattern": "/p", "parser": "js"}],
        "parsers": {"js": {"type": "html", "row_css": ".item", "fields": {"title": {"css": "span.t"}}}},
        "pipelines": [{"type": "filter", "field": "title", "op": "non_empty"}],
        "storage": {"type": "jsonl", "name": name},
        "output": {"dir": "outputs", "base_name": name},
        "anti_bot": {"min_interval": 0.02, "max_retries": 1},
    }
    d = make_task(name, cfg)
    try:
        run_task(d)
        from openpyxl import load_workbook
        wb = load_workbook(str(ROOT / "outputs" / f"{name}.xlsx"))
        ws = wb.active
        check(name, ws.max_row == 5 and ws.max_column >= 1,
              f"rows={ws.max_row} cols={ws.max_column}")
    finally:
        cleanup(name)

    # 3. webhook 通知中间件
    print("[3/10] webhook 中间件批量通知")
    captured = []
    class CapHandler(BaseHTTPRequestHandler):
        def do_POST(self):
            n = int(self.headers.get("Content-Length", 0))
            captured.append(json.loads(self.rfile.read(n).decode()))
            self.send_response(200)
            self.send_header("Content-Length", "0")
            self.end_headers()
        def log_message(self, *a):
            pass
    cap = ThreadingHTTPServer(("127.0.0.1", 0), CapHandler)
    threading.Thread(target=cap.serve_forever, daemon=True).start()
    name = "v3_webhook"
    cfg = {
        "name": name, "start_urls": [f"{base}/p1.html"],
        "queue": {"max_depth": 1, "max_requests": 5, "max_concurrency": 1},
        "source": {"type": "http"},
        "rules": [{"match": "contains", "pattern": "/p", "parser": "js"}],
        "parsers": {"js": {"type": "html", "row_css": ".item", "fields": {"title": {"css": "span.t"}}}},
        "pipelines": [{"type": "filter", "field": "title", "op": "non_empty"}],
        "middleware": [{"on": "data", "action": "webhook",
                        "url": f"http://127.0.0.1:{cap.server_address[1]}/catch", "batch_size": 1}],
        "storage": {"type": "jsonl", "name": name},
        "output": {"dir": "outputs", "base_name": name},
        "anti_bot": {"min_interval": 0.02, "max_retries": 1},
    }
    d = make_task(name, cfg)
    try:
        r = run_task(d)
        titles = [b.get("items", [{}])[0].get("title") for b in captured if b.get("event") == "batch"]
        check(name, len(titles) == 2 and "p1-title-1" in titles,
              f"posts={len(captured)} titles={titles}")
    finally:
        cleanup(name)
        cap.shutdown()

    # 4. jobs 编排
    print("[4/10] jobs 编排（两个任务顺序跑）")
    names = ["v4a", "v4b"]
    for i, nm in enumerate(names):
        cfg = {
            "name": nm, "start_urls": [f"{base}/p{i+1}.html"],
            "queue": {"max_depth": 1, "max_requests": 5, "max_concurrency": 1},
            "source": {"type": "http"},
            "rules": [{"match": "contains", "pattern": "/p", "parser": "js"}],
            "parsers": {"js": {"type": "html", "row_css": ".item", "fields": {"title": {"css": "span.t"}}}},
            "pipelines": [{"type": "filter", "field": "title", "op": "non_empty"}],
            "storage": {"type": "jsonl", "name": nm},
            "output": {"dir": "outputs", "base_name": nm},
            "anti_bot": {"min_interval": 0.02, "max_retries": 1},
        }
        make_task(nm, cfg)
    jobs_fp = ROOT / "outputs" / ".test_jobs.json"
    jobs_fp.write_text(json.dumps([{"task": str(TMP / "v4a")}, {"task": str(TMP / "v4b")}]), encoding="utf-8")
    try:
        p = subprocess.run([PY, "-m", "universal_scraper.cli", "jobs", "--file", str(jobs_fp)],
                           capture_output=True, text=True, cwd=ROOT)
        data = json.loads(p.stdout.strip().splitlines()[-1])
        check("v4_jobs", data.get("jobs") == 2 and data.get("total_items") == 4,
              f"stdout={p.stdout.strip().splitlines()[-1]}")
    finally:
        jobs_fp.unlink(missing_ok=True)
        for nm in names:
            cleanup(nm)

    # 5. TableParser
    print("[5/10] 表格解析（TableParser）")
    name = "v5_table"
    cfg = {
        "name": name, "start_urls": [f"{base}/table.html"],
        "queue": {"max_depth": 1, "max_requests": 5, "max_concurrency": 1},
        "source": {"type": "http"},
        "rules": [{"match": "contains", "pattern": "/table", "parser": "tbl"}],
        "parsers": {"tbl": {"type": "table"}},
        "pipelines": [], "storage": {"type": "jsonl", "name": name},
        "output": {"dir": "outputs", "base_name": name},
        "anti_bot": {"min_interval": 0.02, "max_retries": 1},
    }
    d = make_task(name, cfg)
    try:
        r = run_task(d)
        items = read_items(name)
        ok = r["total"] == 2 and any(i.get("Name") == "Alice" and i.get("Age") == "30" for i in items)
        check(name, ok, f"items={items}")
    finally:
        cleanup(name)

    # 6. LLM 解析打桩
    print("[6/10] LLM 解析（打桩 extract_json）")
    import universal_scraper.llm as llm_mod
    orig = llm_mod.LLMClient.extract_json
    llm_mod.LLMClient.extract_json = lambda self, content, schema, instruction="": {"标题": "AI标题", "价格": "9.9"}
    name = "v6_llm"
    cfg = {
        "name": name, "start_urls": [f"{base}/article.html"],
        "queue": {"max_depth": 1, "max_requests": 5, "max_concurrency": 1},
        "source": {"type": "http"},
        "rules": [{"match": "contains", "pattern": "/article", "parser": "ai"}],
        "parsers": {"ai": {"type": "llm", "schema": {"标题": "文章标题", "价格": "价格"}}},
        "pipelines": [], "storage": {"type": "jsonl", "name": name},
        "output": {"dir": "outputs", "base_name": name},
        "anti_bot": {"min_interval": 0.02, "max_retries": 1},
    }
    d = make_task(name, cfg)
    try:
        r = run_task(d)
        items = read_items(name)
        check(name, r["total"] == 1 and items and items[0].get("标题") == "AI标题",
              f"items={items}")
    finally:
        llm_mod.LLMClient.extract_json = orig
        cleanup(name)

    # 7. 嵌套 jpath
    print("[7/10] 嵌套 JSON 路径（data.list[*] + items[0].x）")
    name = "v7_nested"
    cfg = {
        "name": name, "start_urls": [f"{base}/api/nested"],
        "queue": {"max_depth": 1, "max_requests": 5, "max_concurrency": 1},
        "source": {"type": "http"},
        "rules": [{"match": "contains", "pattern": "/api/nested", "parser": "js"}],
        "parsers": {"js": {"type": "json", "records_path": "data.list",
                           "fields": {"id": {"from": "id"}, "x": {"from": "items[0].x"}}}},
        "pipelines": [], "storage": {"type": "jsonl", "name": name},
        "output": {"dir": "outputs", "base_name": name},
        "anti_bot": {"min_interval": 0.02, "max_retries": 1},
    }
    d = make_task(name, cfg)
    try:
        r = run_task(d)
        items = read_items(name)
        check(name, r["total"] == 2 and {i["x"] for i in items} == {"a", "b"},
              f"items={items}")
    finally:
        cleanup(name)

    # 8. deny + same_domain 组合
    print("[8/10] extract_links deny + same_domain 组合")
    name = "v8_deny"
    cfg = {
        "name": name, "start_urls": [f"{base}/hubdocs.html"],
        "queue": {"max_depth": 2, "max_requests": 10, "max_concurrency": 2},
        "source": {"type": "http"},
        "rules": [{"match": "contains", "pattern": "/hubdocs", "parser": "list"},
                  {"match": "contains", "pattern": "/docs", "parser": "js"},
                  {"match": "contains", "pattern": "/blog", "parser": "js"}],
        "parsers": {"list": {"type": "html", "fields": {"h": {"css": "h1"}},
                             "extract_links": {"same_domain": True, "deny": "/blog"}},
                    "js": {"type": "html", "row_css": ".item", "fields": {"t": {"css": "span.t"}}}},
        "pipelines": [{"type": "filter", "field": "h", "op": "non_empty"}],
        "storage": {"type": "jsonl", "name": name},
        "output": {"dir": "outputs", "base_name": name},
        "anti_bot": {"min_interval": 0.02, "max_retries": 1},
    }
    d = make_task(name, cfg)
    try:
        r = run_task(d)
        urls = [i["_url"] for i in read_items(name)]
        check(name, r["fetched"] == 2 and not any("blog" in u for u in urls),
              f"fetched={r['fetched']} urls={urls}")
    finally:
        cleanup(name)

    # 9. 配置校验拦截
    print("[9/10] 配置校验拦截坏配置")
    bad1 = {"name": "x", "start_urls": ["http://a"], "source": {"type": "browser", "actions": [{"type": "explode"}]},
            "rules": [{"match": "contains", "pattern": "/", "parser": "p"}], "parsers": {"p": {"type": "html"}},
            "storage": {"type": "jsonl"}}
    bad2 = {"name": "x", "start_urls": ["http://a"], "source": {"type": "http"},
            "rules": [{"match": "contains", "pattern": "/", "parser": "p"}], "parsers": {"p": {"type": "html"}},
            "storage": {"type": "nosql"}}
    bad3 = {"name": "x", "start_urls": ["http://a"], "source": {"type": "http"},
            "rules": [{"match": "contains", "pattern": "/", "parser": "missing"}], "parsers": {"p": {"type": "html"}},
            "storage": {"type": "jsonl"}}
    ok1 = ok2 = ok3 = False
    try:
        validate_task(bad1)
    except ConfigError:
        ok1 = True
    try:
        validate_task(bad2)
    except ConfigError:
        ok2 = True
    try:
        validate_task(bad3)
    except ConfigError:
        ok3 = True
    check("v9_validate", ok1 and ok2 and ok3, f"bad_action={ok1} bad_storage={ok2} bad_parser={ok3}")

    # 10. fetch --browser --selector 组合
    print("[10/10] fetch --browser --selector（JS 渲染页）")
    p = subprocess.run([PY, "-m", "universal_scraper.cli", "fetch", f"{base}/spa.html",
                        "--browser", "--selector", ".item"], capture_output=True, text=True, cwd=ROOT)
    ok = "spa-item-1" in p.stdout and "spa-item-3" in p.stdout
    check("v10_fetch_browser_selector", ok, f"stdout={p.stdout[:120]}")

    print(f"\n===== 结果: {len(PASS)}/10 通过 =====")
    if FAIL:
        print("失败:", FAIL)
        srv.shutdown()
        return 1
    srv.shutdown()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
