#!/usr/bin/env python3
"""📋 名单类附件解析（通报/公示/招标清单：一份附件 = 一张名单）。

与 `pdf_table.py` 的分工：pdf_table 是**通用单表抽取**（PDF/xlsx → 行/文本）；
本模块是**名单语义层**——跨期列名对齐、行语义合并、声明数对账，并复用 pdf_table 的
xlsx 解析（`.xlsx` 链），PDF/docx/doc 用自己的名单版装配（pdf_table 的通用抽表
不处理 rowspan 续行/一问题一行这些名单形态）。

五条解析链（`parse_attachment_ex` 报告实际用了哪条 + 失败原因）：
  ① docx  → python-docx 表格
  ② doc   → LibreOffice 转 HTML（真表格标记，列位可靠）；textutil 扁平分格兜底
  ③ pdf   → pdfplumber 抽表（**跨页串接**后统一装配）；无文本层则渲染成图走 ④
  ④ 图片  → 云端 VLM 逐行 + 本地 Vision 交叉核验
  ⑤ 内联  → 正文 HTML 表格（部分通报把名单直接贴正文）
  （另有 .xlsx/.xls 链复用 pdf_table）

行语义（全部是实战踩出来的，逐条对应一个真实缺陷）：
  - 表头驱动列位映射：**空单元格不能删**，否则整行左移（"企业名称"被读成"应用名称"）
  - rowspan 续行并回：前几列 rowspan=N 时，每个物理行只装一格（41 个 APP 被读成 89 行）
  - "一问题一行"合并：同一 APP 按所涉问题拆成多行（89 行序号只到 60）
  - 同 APP 多来源/版本合并：序号+名称相同（1 款 APP 列了 App Store 与安卓两个来源）
  - 无序号即续行：**仅当该段序号列至少出现 2 个序号值**（Word 自动编号不落文本时整列为空）
  - 换行碎片并回：单元格折行被切成独立行（"上海邮乐网络"/"技术有限公司"）
  - 合计/小计/备注行、分节行：跳过并计入 stats["dropped"]（并进上一条会把"合计 41款"
    写成上一条 APP 的"所涉问题"）

声明数对账：正文"发现N款/尚有N款"是**句子级求和**（部本级 + 各省管理局分列），
且必须排除累计口径句（"对…368款APP提出整改要求"）与汇总复述句（"对上述共计106款进行下架"）。

诊断纪律（0 结果必须可归因）：`aggregate_ex` 逐附件回传 chain/rows/error，
0 行与失败也进账；`audit_counts_ex` 把"没有声明数可对账"单列成 `no_stated`，
不让"根本没做对账"伪装成"全部达标"。
"""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

F_NAME = ("软件名称", "APP名称", "应用名称", "App名称", "应用名", "名称", "应用软件名称", "产品名称")
F_COMP = ("企业名称", "开发者", "应用开发者", "公司名称", "开发者名称", "运营者名称", "生产厂商")
F_VER = ("版本", "版本号", "应用版本", "复测应用版本", "复测版本", "下架版本")
F_SRC = ("版本来源", "应用来源", "应用来源商店", "来源", "应用商店", "来源商店", "样品来源")
F_ISSUE = ("所涉问题", "涉及问题", "问题", "问题项")
FIELD_SETS = (("app_name", F_NAME), ("company_name", F_COMP), ("version", F_VER),
              ("store_source", F_SRC), ("issue_types", F_ISSUE))
CELL = "\x07"
VER_LIKE = re.compile(r"[vV]?(?:[A-Za-z]+[_.\-])*\d[A-Za-z0-9.\-_]*")   # 不含 CJK
VER_LIKE_LOOSE = re.compile(r"[vV]?\d+(\.\d+)+[A-Za-z0-9.\-_]*")   # 兜底链的宽松形（同样不含 CJK）
SECTION_RE = re.compile(r"^[（(][一二三四五六七八九十]+[)）]")
TOTAL_ROW_RE = re.compile(r"^(合计|总计|小计|以上|以下|备注|说明|注[:：]|资料来源|数据来源)")
COMPANY_STRONG = ("有限公司", "有限责任公司", "公司", "集团", "中心", "研究院", "研究所", "大学",
                  "学院", "银行", "医院", "工作室", "协会", "报社", "电视台", "出版社", "事业部",
                  "管理局", "总局", "分行", "事务所", "连锁", "药房", "影院", "酒店", "网络科技",
                  "信息技术", "传媒")
STORE_RE = re.compile(r"(官网|商店|市场|应用宝|豌豆荚|百度|华为|小米|oppo|vivo|360|应用汇|搜狗|"
                      r"App ?Store|网站|网$|酷安|应用中心)", re.I)
ISSUE_HINT = re.compile(r"(收集|权限|推送|注销|弹窗|强制|欺骗|误导|骚扰|索取|超范围|私自|共享|泄露|"
                        r"SDK|广告|捆绑|不给|定向|账号|隐私|通讯录|麦克风|相册|位置|通讯|关闭|"
                        r"跳转|自动|违规|问题|收集使用|两次|默认)")
ROW_PROMPT = ("这是名单表格中的一行（单元格可能折成两三行）。请逐字提取这一行的字段，"
              "输出 JSON 对象，键固定为：序号、应用名称、应用开发者、应用来源、应用版本、所涉问题。"
              "所涉问题若有多个用 ； 分隔。看不清的字段留空字符串，不要猜测、不要补全。只输出 JSON。")

HEADER_NS = {re.sub(r"\s+", "", w) for w in list(F_NAME + F_COMP + F_VER + F_SRC + F_ISSUE)
             + ["序号", "应用类别", "类别", "复测版本"]}


# ---------------- 文本工具 ----------------

def _norm(s: Any) -> str:
    return re.sub(r"\s+", " ", str(s or "").replace("\u3000", " ")).strip()


def _clean(s: Any) -> str:
    """保留换行的清洗（问题列多行 = 多个问题，不能压成一行）。"""
    s = str(s or "").replace("\u3000", " ").replace("\xa0", " ")
    s = re.sub(r"[ \t\r\f\v]+", " ", s)
    return "\n".join(x.strip() for x in s.split("\n")).strip()


def _norm_ws(s: Any) -> str:
    """去全部空白——表头常写成"应用\\n来源""应用 版本"。"""
    return re.sub(r"\s+", "", str(s or "").replace("\u3000", ""))


def _split_issues(cell: Any) -> List[str]:
    parts = [p.strip(" ；;、,，") for p in re.split(r"[\n;；]+", str(cell or ""))]
    return [p for p in parts if p and _norm_ws(p) not in HEADER_NS]


def _is_seq(c: str) -> bool:
    return bool(re.fullmatch(r"\d{1,4}", _norm(c)))


def _is_company(c: str) -> bool:
    c = _norm(c)
    return bool(c) and len(c) <= 40 and any(w in c for w in COMPANY_STRONG)


def _is_store(c: str) -> bool:
    c = _norm(c)
    return bool(c) and bool(STORE_RE.search(c)) and len(c) <= 20


def _looks_issue(c: str) -> bool:
    c = _norm(c)
    return bool(ISSUE_HINT.search(c)) and len(c) <= 60


def _is_section(c: str) -> bool:
    c = _norm(c)
    return bool(SECTION_RE.match(c)) or (c.endswith("名单") and len(c) <= 30)


# ---------------- 表头与行装配 ----------------

def _field_of(cell: str) -> Optional[str]:
    w = _norm_ws(cell)
    if w == "序号":
        return "seq"
    for fld, words in FIELD_SETS:
        if w in {_norm_ws(x) for x in words}:
            return fld
    return None


def header_map(cells: Sequence[str]) -> Optional[Dict[str, int]]:
    """表头 → {字段: 列号}；识别不出（<3 列或没有名称列）返回 None。"""
    m: Dict[str, int] = {}
    for i, c in enumerate(cells):
        fld = _field_of(c)
        if fld:
            m.setdefault(fld, i)
    return m if "app_name" in m and len(m) >= 3 else None


def is_header_row(cells: Sequence[str]) -> bool:
    """表头行判定：允许一个未知格（别名滞后一次就漏判——实测"应用版本"漏判后
    序号列可靠性预判失效，换行碎片全部漏并）。"""
    non = [c for c in cells if _norm_ws(c)]
    if not non:
        return False
    hit = sum(1 for c in non if _norm_ws(c) in HEADER_NS)
    return hit >= max(2, len(non) - 1)


def _row_by_header(cells: Sequence[str], hm: Dict[str, int]) -> Optional[Dict[str, Any]]:
    r: Dict[str, Any] = {}
    for fld, i in hm.items():
        v = cells[i] if i < len(cells) else ""
        if fld == "seq":
            if _is_seq(v):
                r["seq"] = int(_norm(v))
        elif fld == "issue_types":
            r["issue_types"] = _split_issues(v)
        else:
            r[fld] = _norm(v)
    if not r.get("app_name"):
        return None
    r.setdefault("issue_types", [])
    return r


def _row_from_cells(cells: Sequence[str]) -> Optional[Dict[str, Any]]:
    """无表头兜底：序号取前三格内首个整数（兼容"应用类别"前置列），余下按值形状。"""
    plate = [_norm(c) for c in cells if _norm(c)]
    if not plate or is_header_row(plate) or all(_is_seq(c) for c in plate):
        return None
    si = next((k for k in range(min(3, len(plate))) if _is_seq(plate[k])), None)
    i = si + 1 if si is not None else 0
    rest = plate[i:]
    if len(rest) < 2:
        return None
    r: Dict[str, Any] = {"app_name": rest[0], "company_name": rest[1]}
    if si is not None:
        r["seq"] = int(plate[si])
    for c in rest[2:]:
        if "version" not in r and VER_LIKE.fullmatch(c):
            r["version"] = c
        elif "store_source" not in r and _is_store(c):
            r["store_source"] = c
        else:
            r.setdefault("issue_types", []).extend(_split_issues(c))
    r.setdefault("issue_types", [])
    return r


def _append_issues(prev: Dict[str, Any], v: str) -> None:
    lst = prev.setdefault("issue_types", [])
    for it in _split_issues(v):
        if it not in lst:
            lst.append(it)


def _merge_fragment(prev: Dict[str, Any], cells: Sequence[str], hm: Dict[str, int]) -> int:
    """换行碎片并回上一条记录（按列位；内容形状可覆盖列位）。

    返回**实际并入了几个字段**：0 表示这一行的内容全落在未映射列（备注/类别…）或
    序号列——调用方必须据此落回兜底装配，不能把行直接吞掉（实测"备注"列的续行
    文本会凭空消失）。
    """
    merged = 0
    for fld, i in sorted(hm.items(), key=lambda kv: kv[1]):
        v = _norm(cells[i] if i < len(cells) else "")
        if not v or fld == "seq":
            continue
        if fld == "issue_types" or _looks_issue(v):
            _append_issues(prev, v)
            merged += 1
            continue
        if fld in ("app_name", "version") and VER_LIKE.fullmatch(v) and not prev.get("version"):
            prev["version"] = v
            merged += 1
            continue
        if fld in ("app_name", "company_name", "store_source"):
            prev[fld] = (prev.get(fld, "") + v).strip()      # 换行处无空格，直连
            merged += 1
        else:
            prev[fld] = (prev.get(fld, "") + v).strip()
            merged += 1
    return merged


def _is_continuation_row(cells: Sequence[str], hm: Dict[str, int]) -> bool:
    """rowspan 续行判定：① 非空格全落在"所涉问题"列及其右侧 ② 整行只有一格。

    单格行还要排除表头词与 注/说明/备注/合计/总计… 前缀——它们是表格脚注，
    并进上一条会把"合计 41款"写成上一条 APP 的"所涉问题"（实测）。"""
    idx = [i for i, c in enumerate(cells) if _norm(c)]
    if not idx or hm.get("issue_types") is None:
        return False
    if len(idx) == 1:
        v = _norm(cells[idx[0]])
        if _norm_ws(v) in HEADER_NS or TOTAL_ROW_RE.match(v):
            return False
        return True
    return min(idx) >= hm["issue_types"]


def _table_segments(cells_rows: Sequence[Sequence[str]]) -> List[Tuple[int, Dict[str, int], bool]]:
    """按表头切段 → [(行号, 列映射, 该段序号列是否确有值)]。

    为什么必须**分段**：`parse_docx`/`parse_pdf`/内联都把整篇表格串接成一份行序列，
    后段表头可能没有序号列（或序号整列为空）。`seq_ok` 只在首段算一次的话，
    后段每一行都会被"无序号即续行"吞掉（实测 6 行塌成 2 行、名称粘连）。
    """
    segs: List[List[Any]] = []
    for k, cells in enumerate(cells_rows):
        plate = [c for c in cells if _norm(c)]
        if not plate:
            continue
        hm = header_map(cells)
        if not hm:
            continue
        # 数据行偶尔也会凑出 ≥3 个字段名（如名称就叫"版本"），用"有数字序号格"排除
        if not is_header_row(plate) and any(_is_seq(c) for c in cells):
            continue
        segs.append([k, hm, False])
    for i, seg in enumerate(segs):
        k, hm = seg[0], seg[1]
        end = segs[i + 1][0] if i + 1 < len(segs) else len(cells_rows)
        if "seq" in hm:
            vals = []
            for cells in cells_rows[k + 1:end]:
                plate = [c for c in cells if _norm(c)]
                if not plate or is_header_row(plate):
                    continue
                if hm["seq"] < len(cells):
                    vals.append(_is_seq(cells[hm["seq"]]))
            seg[2] = sum(vals) >= 2
    return [(s[0], s[1], s[2]) for s in segs]


def table_rows(cells_rows: Sequence[Sequence[str]],
               stats: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
    """表格行装配（整篇串接后调用：跨页续表往往没有表头，靠首表列映射）。

    列位对齐必须用**原始单元格位置**；"无序号即续行"仅在**该段**序号列确有值时启用
    （≥2 个序号值；序号格整列为空时误用会把 41 行并成 1 行）。

    传 `stats`（dict）时回填诊断：`segments` / `merged_fragments` / `dropped`
    （丢行样本 [(行号, 原因, 前几格)]）——0 行必须可归因，不能只有"空结果"。
    """
    segs = _table_segments(cells_rows)
    by_idx = {k: (hm, ok) for k, hm, ok in segs}
    if stats is not None:
        stats.setdefault("segments", len(segs))
        stats.setdefault("merged_fragments", 0)
        stats.setdefault("dropped", [])
        stats.setdefault("skipped_section", 0)

    def _drop(idx: int, reason: str, cells: Sequence[str]) -> None:
        if stats is not None:
            stats["dropped"].append((idx, reason, [_norm(c) for c in cells if _norm(c)][:3]))

    hm: Optional[Dict[str, int]] = None
    seq_ok = False
    out: List[Dict[str, Any]] = []
    for idx, cells in enumerate(cells_rows):
        plate = [c for c in cells if _norm(c)]
        if not plate:
            continue
        if idx in by_idx:                                  # 段首表头：整段换映射
            hm, seq_ok = by_idx[idx]
            continue
        if is_header_row(plate):                           # 段中重复表头（跨页表头行）
            continue
        if len(plate) >= 2 and len({_norm_ws(c) for c in plate}) == 1:
            _drop(idx, "merged_cell_section", cells)       # 合并单元格的分节行
            continue
        if len(plate) == 1 and _is_section(plate[0]):
            if stats is not None:
                stats["skipped_section"] += 1
            _drop(idx, "section", cells)
            continue
        if TOTAL_ROW_RE.match(plate[0]) and (len(plate) == 1 or not _is_seq(plate[0])):
            _drop(idx, "total_row", cells)                 # 合计/小计/备注行
            continue
        row_seq = None
        if hm and "seq" in hm and hm["seq"] < len(cells):
            row_seq = int(_norm(cells[hm["seq"]])) if _is_seq(cells[hm["seq"]]) else None
        if out and hm and _is_continuation_row(cells, hm):
            n = _merge_fragment(out[-1], cells, hm)
            if n == 0:
                _drop(idx, "fragment_unmapped", cells)     # 未映射列的续行：落兜底，别吞
                r = _row_from_cells(plate)
                if r:
                    _lift_version(r)
                    out.append(r)
            elif stats is not None:
                stats["merged_fragments"] += 1
            continue
        if out and hm and seq_ok and row_seq is None:
            n = _merge_fragment(out[-1], cells, hm)
            if n == 0:
                _drop(idx, "fragment_unmapped", cells)
                r = _row_from_cells(plate)
                if r:
                    _lift_version(r)
                    out.append(r)
            elif stats is not None:
                stats["merged_fragments"] += 1
            continue
        r = _row_by_header(cells, hm) if hm else None
        if r is None:
            r = _row_from_cells(plate)
        if r:
            _lift_version(r)
            out.append(r)
        else:
            _drop(idx, "unassemblable", cells)
    return out


def _lift_version(r: Dict[str, Any]) -> None:
    """问题列里的版本形条目归位到 version（PDF 换行把"4.5.17.7"拆进问题列）。"""
    if r.get("version"):
        return
    lst = r.get("issue_types") or []
    vers = [x for x in lst if VER_LIKE.fullmatch(_norm(x)) and re.search(r"\d", x)]
    if vers and len(vers) <= 2 and all(len(x) <= 12 for x in vers):
        r["version"] = "".join(_norm(x) for x in vers)
        r["issue_types"] = [x for x in lst if x not in vers]


def merge_app_rows(rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """一 APP 一行：① 身份字段全同（一问题一行）② 序号+名称相同（多来源/版本，
    额外来源进 `alt_sources`）。"""
    out: List[Dict[str, Any]] = []
    for r in rows:
        if out:
            p = out[-1]
            same = all(_norm(p.get(k)) == _norm(r.get(k))
                       for k in ("app_name", "company_name", "version", "store_source"))
            same_app = (r.get("seq") is not None and p.get("seq") == r.get("seq")
                        and _norm(p.get("app_name")) and _norm(p.get("app_name")) == _norm(r.get("app_name")))
            if (same or same_app) and (p.get("app_name") or p.get("company_name")):
                if same_app and not same:
                    alt = {k: r.get(k) for k in ("company_name", "version", "store_source")}
                    if any(alt.values()):
                        p.setdefault("alt_sources", []).append(alt)
                for it in (r.get("issue_types") or []):
                    if it not in p.setdefault("issue_types", []):
                        p["issue_types"].append(it)
                continue
        out.append(dict(r))
    return out


# ---------------- 五条解析链 ----------------

def _doc_html_ex(path: Path, cache_dir: Optional[Path] = None) -> Tuple[str, str]:
    """LibreOffice → HTML；返回 (html, error)。缓存键含父目录名（附件常同名）。"""
    from .core import run_tool
    cache = (Path(cache_dir) if cache_dir else Path("out/doc_html")) / f"{path.parent.name}__{path.stem}.html"
    if cache.exists() and cache.stat().st_size > 0:
        return cache.read_text(encoding="utf-8", errors="replace"), ""
    cache.parent.mkdir(parents=True, exist_ok=True)
    r = run_tool(["soffice", "--headless", "--convert-to", "html",
                  "--outdir", str(cache.parent), str(path)])
    if not r["ok"]:
        # 失败必须带工具名/退出码/stderr 首段——"两条 doc 链都没出结果"这种话无法定位
        return "", f"soffice code={r['code']} {r['err'].strip()[:120]}"
    produced = cache.parent / (path.stem + ".html")
    if produced.exists() and produced.stat().st_size > 0:
        produced.replace(cache)
    if cache.exists():
        return cache.read_text(encoding="utf-8", errors="replace"), ""
    return "", "soffice 未产出 HTML（转换静默失败）"


def _doc_html(path: Path, cache_dir: Optional[Path] = None) -> str:
    """LibreOffice → HTML（真 <table>）。失败返回空串（要原因用 _doc_html_ex）。"""
    return _doc_html_ex(path, cache_dir)[0]


def _textutil_cells_ex(path: Path) -> Tuple[List[str], str]:
    """textutil 兜底：多数 .doc 用 \\x07 分隔单元格；文本型 doc 只有换行。"""
    from .core import run_tool
    r = run_tool(["textutil", "-convert", "txt", "-stdout", str(path)])
    if not r["ok"]:
        return [], f"textutil code={r['code']} {r['err'].strip()[:120]}"
    if not r["out"]:
        return [], "textutil 输出为空"
    return [_norm(c) for c in re.split(f"[{CELL}\n]", r["out"])], ""


def _textutil_cells(path: Path) -> List[str]:
    return _textutil_cells_ex(path)[0]


def _assemble_doc_rows(cells: Sequence[str]) -> List[Dict[str, Any]]:
    """textutil 扁平流装配（无真表格时的兜底）：结构恒定 名称→企业→版本→来源→问题…"""
    cells = [c for c in cells if c and _norm_ws(c) not in HEADER_NS
             and not c.startswith("附件") and not _is_section(c)]
    starts: List[int] = []
    n = len(cells)
    for i, c in enumerate(cells):
        nxt = cells[i + 1] if i + 1 < n else ""
        if _is_company(nxt) and not _is_company(c) and not _is_seq(c) and not _looks_issue(c):
            starts.append(i)
        elif _is_seq(c) and nxt and not _is_seq(nxt) and not _is_store(nxt) \
                and not _looks_issue(nxt) and not _is_company(nxt):
            starts.append(i)
        elif nxt and c == nxt and 1 < len(c) <= 30 and not _is_company(c) and not _looks_issue(c):
            starts.append(i)
    have = set(starts)
    for j, c in enumerate(cells):
        if not VER_LIKE_LOOSE.fullmatch(c) or j < 2:
            continue
        if j - 2 in have or j - 1 in have or (j >= 3 and j - 3 in have):
            continue
        cand = cells[j - 2]
        if cand and not _is_store(cand) and not _looks_issue(cand) and not _is_seq(cand) \
                and not _is_company(cand):
            starts.append(j - 2)
    merged: List[int] = []
    for s in sorted(set(starts)):
        if merged and s == merged[-1] + 1:
            continue
        merged.append(s)
    rows: List[Dict[str, Any]] = []
    for k, s in enumerate(merged):
        e = merged[k + 1] if k + 1 < len(merged) else n
        seg = [c for c in cells[s:e] if c]
        if len(seg) < 2:
            continue
        r: Dict[str, Any] = {}
        i = 0
        if _is_seq(seg[0]):
            r["seq"] = int(_norm(seg[0]))
            i = 1
        if i >= len(seg):
            continue
        r["app_name"] = seg[i]
        i += 1
        if i < len(seg) and _is_company(seg[i]):
            r["company_name"] = seg[i]
            i += 1
        elif i < len(seg) and seg[i] == r["app_name"]:
            i += 1                                    # 合并单元格：企业格为空致名称重复
        for c in seg[i:]:
            if "version" not in r and VER_LIKE.fullmatch(c):
                r["version"] = c
            elif "store_source" not in r and _is_store(c):
                r["store_source"] = c
            else:
                r.setdefault("issue_types", []).extend(_split_issues(c))
        if r.get("app_name"):
            rows.append(r)
    return rows


def _html_grid(table) -> List[List[str]]:
    """按 rowspan/colspan 还原完整网格（rowspan 会让后续物理行整体左移）。"""
    grid: List[List[str]] = []
    pending: Dict[int, List[Any]] = {}
    for tr in table.xpath(".//tr"):
        row: List[str] = []
        col = 0
        cells = tr.xpath("./td|./th")
        ci = 0
        guard = 0
        while ci < len(cells) or pending:
            guard += 1
            if guard > 500:
                break
            if col in pending:
                remain, txt = pending[col]
                row.append(txt)
                if remain <= 1:
                    del pending[col]
                else:
                    pending[col] = [remain - 1, txt]
                col += 1
                continue
            if ci >= len(cells):
                row.append("")
                col += 1
                continue
            td = cells[ci]
            ci += 1
            txt = _cell_html_text(td)
            rs = int(td.get("rowspan") or 1)
            cs = int(td.get("colspan") or 1)
            for _ in range(cs):
                row.append(txt)
                if rs > 1:
                    pending[col] = [rs - 1, txt]
                col += 1
        grid.append(row)
    return grid


def _cell_html_text(td) -> str:
    """单元格文本：<p>/<br> 边界转 \\n（问题列多行才不会被压成一行）。"""
    parts: List[str] = []
    for node in td.iter():
        if node.tag in ("p", "br", "div", "li", "tr"):
            parts.append("\n")
        if node.text:
            parts.append(node.text)
        if node.tail:
            parts.append(node.tail)
    return _clean("".join(parts))


def parse_inline_html(html: str) -> List[Dict[str, Any]]:
    """正文内联表格（多张表整篇串接：第二张常无表头）。"""
    try:
        from lxml import html as LH
        doc = LH.fromstring(html)
    except Exception:
        return []
    rows: List[List[str]] = []
    for table in doc.xpath("//table"):
        rows += _html_grid(table)
    return table_rows(rows)


def parse_docx(path: Path) -> List[Dict[str, Any]]:
    try:
        import docx
    except ImportError:
        return _assemble_doc_rows(_textutil_cells(path))
    try:
        d = docx.Document(str(path))
    except Exception:
        return []
    rows: List[List[str]] = []
    for table in d.tables:
        rows += [[_clean(c.text) for c in tr.cells] for tr in table.rows]
    return table_rows(rows)


def parse_doc(path: Path, cache_dir: Optional[Path] = None) -> List[Dict[str, Any]]:
    """.doc：LibreOffice→HTML 主链；textutil 兜底。"""
    html = _doc_html(path, cache_dir)
    if html:
        rows = parse_inline_html(html)
        if rows:
            return rows
    return _assemble_doc_rows(_textutil_cells(path))


def parse_pdf(path: Path) -> List[Dict[str, Any]]:
    try:
        import pdfplumber
    except ImportError:
        return []
    rows: List[List[str]] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                for table in (page.extract_tables() or []):
                    rows += [[_clean(c) for c in tr] for tr in table]
    except Exception:
        return []
    return table_rows(rows)


def _scanned_pdf_pngs_ex(path: Path, out_dir: Path, resolution: int = 170) -> Tuple[List[Path], str]:
    """扫描版 PDF（无文本层）逐页渲染 PNG；返回 (pngs, error)。

    一页渲染失败**保留已渲染的页**（旧实现直接丢整份，已成功的前几页白渲染），
    并把异常类型与页码带出来。"""
    try:
        import pdfplumber
    except ImportError:
        return [], "pdfplumber 未安装"
    out_dir.mkdir(parents=True, exist_ok=True)
    pngs: List[Path] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            if any((pg.extract_text() or "").strip() for pg in pdf.pages):
                return [], "有文本层（不是扫描件；表格抽不出属版式问题，非 OCR 问题）"
            for i, pg in enumerate(pdf.pages, 1):
                p = out_dir / f"{path.parent.name[-8:]}_{path.stem}_p{i}.png"
                if not p.exists():
                    pg.to_image(resolution=resolution).save(str(p))
                pngs.append(p)
    except Exception as e:
        return pngs, f"渲染中断于第 {len(pngs) + 1} 页: {type(e).__name__}: {e}"
    return pngs, ""


def _scanned_pdf_pngs(path: Path, out_dir: Path, resolution: int = 170) -> List[Path]:
    return _scanned_pdf_pngs_ex(path, out_dir, resolution)[0]


def _vision_obj(llm: Any, prompt: str, image: Any) -> Dict[str, Any]:
    """单次 VLM 调用 → dict（容忍 ```json 包裹与解释文字）。"""
    import base64
    import io
    if isinstance(image, Path):
        b64 = base64.b64encode(Path(image).read_bytes()).decode()
    else:
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        b64 = base64.b64encode(buf.getvalue()).decode()
    raw = llm.vision(prompt, f"data:image/png;base64,{b64}")
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", (raw or "").strip())
    m = re.search(r"\{.*\}", raw, re.S)
    if not m:
        return {}
    try:
        o = json.loads(m.group(0))
        return o if isinstance(o, dict) else {}
    except Exception:
        return {}


LIST_PROMPT = ("这是名单表格的截图。请逐行提取表格内容，输出 JSON 数组，每行一个对象，"
               "键固定为：序号、应用名称、开发者、应用来源、版本、所涉问题。"
               "所涉问题若含多个用 ； 分隔。只输出 JSON，不要解释。")


def parse_images(paths: Sequence[Path], llm: Any = None, use_local_ocr: bool = True,
                 ocr_cache: Optional[Path] = None,
                 stats: Optional[Dict[str, Any]] = None) -> List[Dict[str, Any]]:
    """图片名单：**序号列锚定切行** → 逐行 3× 放大 → VLM；本地 Vision 交叉核验公司名。

    整图一次性 OCR 实测公司名错字 ~11%（小图每行只有几十像素高）；逐行放大降到个位数，
    且交叉核验能抓住云端模型的补字偏置（`local_ocr.cross_check_verdict` 判 alt_wins）。

    两条纪律：
    - **本地引擎不可用时不写 `ocr_check`**（写 `unavailable`）——不能把"第二引擎没跑"
      伪装成"两引擎分歧"；
    - **空结果不写缓存**——否则一次瞬时失败（限流/空响应）会被缓存永久固化成
      "这张图没有名单"，重跑也修不回来（实测）。命中空缓存视为未命中，重试。
    传 `stats` 时回填 `failed_bands`（VLM 空返回的带数）与 `ocr_unavailable`（原因）。
    """
    from . import local_ocr
    if llm is None:
        from .llm import LLMClient
        llm = LLMClient()
    # 行切分只需 Pillow；Vision 只用于**交叉核验**——两者不可混为一谈：
    # 引擎不可用时仍要逐行（质量远好于整图），只是不写 ocr_check 的引擎结论。
    local_ok = use_local_ocr and local_ocr.available()
    if use_local_ocr and not local_ok and stats is not None:
        stats["ocr_unavailable"] = local_ocr.unavailable_reason()
    rows: List[Dict[str, Any]] = []
    for p in paths:
        bands = local_ocr.row_bands(p) if use_local_ocr else []
        if not bands:
            rows += _parse_image_whole(p, llm)
            continue
        try:
            from PIL import Image
            img = Image.open(p)
            W, _H = img.size
        except Exception:
            continue
        for (y0, y1) in bands:
            o = None
            cp = None
            if ocr_cache is not None:
                ck = hashlib.sha256(f"{p}|{y0}|{y1}".encode()).hexdigest()[:20]
                cp = Path(ocr_cache) / f"{ck}.json"
                if cp.exists():
                    try:
                        cached = json.loads(cp.read_text(encoding="utf-8"))
                    except Exception:
                        cached = {}
                    if cached:                            # 空缓存 = 上次失败，重试
                        o = cached
                    else:
                        cp.unlink(missing_ok=True)
            if o is None:
                crop = img.crop((0, y0, W, y1))
                crop = crop.resize((crop.size[0] * 3, crop.size[1] * 3), Image.LANCZOS)
                o = _vision_obj(llm, ROW_PROMPT, crop)
                if o and ocr_cache is not None and cp is not None:
                    cp.parent.mkdir(parents=True, exist_ok=True)
                    cp.write_text(json.dumps(o, ensure_ascii=False), encoding="utf-8")
            if not o:
                if stats is not None:
                    stats["failed_bands"] = stats.get("failed_bands", 0) + 1
                continue
            name = _norm(o.get("应用名称"))
            if not name or _norm_ws(name) in HEADER_NS or "序号" in name:
                continue
            r = _row_from_vision(o)
            if not _is_seq(_norm(o.get("序号", ""))):
                if rows:                                     # 折行续带：并入上一行
                    prev = rows[-1]
                    for k in ("app_name", "company_name"):
                        v = r.get(k) or ""
                        if not v:
                            continue
                        if k == "app_name" and (_looks_issue(v) or len(v) > 12):
                            _append_issues(prev, v)
                        else:
                            prev[k] = (prev.get(k, "") + v).strip()
                    if r.get("version") and not prev.get("version"):
                        prev["version"] = r["version"]
                    for it in r.get("issue_types") or []:
                        if it not in prev.setdefault("issue_types", []):
                            prev["issue_types"].append(it)
                    continue
            else:
                r["seq"] = int(_norm(o["序号"]))
            if local_ok:
                vis = _vision_company_in_band(p, y0, y1)
                verdict = local_ocr.cross_check_verdict(r.get("company_name", ""), vis)
                r["ocr_check"] = verdict
                if verdict == "alt_wins":
                    r["company_name"] = vis
                    r["ocr_alt"] = vis
                elif verdict == "differ":
                    r["ocr_alt"] = vis                    # 副读留档（人工复核时能对照）
            elif use_local_ocr:
                r["ocr_check"] = "unavailable"            # 引擎没跑 ≠ 两引擎分歧
            rows.append(r)
    return rows


def _row_from_vision(o: Dict[str, Any]) -> Dict[str, Any]:
    return {"app_name": _norm(o.get("应用名称") or o.get("APP名称")),
            "company_name": _norm(o.get("应用开发者") or o.get("企业名称")),
            "version": _norm(o.get("应用版本") or o.get("版本号")),
            "store_source": _norm(o.get("应用来源") or o.get("版本来源")),
            "issue_types": _split_issues(o.get("所涉问题", ""))}


def _parse_image_whole(path: Path, llm: Any) -> List[Dict[str, Any]]:
    """无框线图的兜底：整图 + 数组提示词（质量低于逐行，仅在没有行带时用）。"""
    import base64
    try:
        b64 = base64.b64encode(Path(path).read_bytes()).decode()
        raw = llm.vision(LIST_PROMPT, f"data:image/png;base64,{b64}")
    except Exception:
        return []
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", (raw or "").strip())
    m = re.search(r"\[.*\]", raw, re.S)
    if not m:
        return []
    try:
        arr = json.loads(m.group(0))
    except Exception:
        return []
    out: List[Dict[str, Any]] = []
    for o in arr if isinstance(arr, list) else []:
        if not isinstance(o, dict):
            continue
        r = _row_from_vision(o)
        if not r["app_name"]:
            continue
        if _is_seq(_norm(o.get("序号", ""))):
            r["seq"] = int(_norm(o["序号"]))
        out.append(r)
    return out


def _vision_company_in_band(path: Path, y0: int, y1: int) -> str:
    """本地 Vision 读该行带内"开发者"列（x≈0.19~0.40）的文本。"""
    from . import local_ocr
    try:
        from PIL import Image
        _W, H = Image.open(path).size
    except Exception:
        return ""
    parts = []
    for s in local_ocr.ocr_image(path):
        y_px = (1 - s["y"]) * H
        if y0 - 4 <= y_px <= y1 + 4 and 0.19 <= s["x"] <= 0.40:
            parts.append(s["text"])
    return "".join(parts).strip()


def _rows_from_keyed_dicts(items: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """把"表头→值"的字典行映射到本模块字段（`pdf_table.parse_xlsx` 的输出形态）。

    该函数的 rows 已经是 `{列名: 值}`（不是单元格列表），所以按**列名**映射，
    不能再喂给 `table_rows`（那是按列位的装配）。"""
    out: List[Dict[str, Any]] = []
    for o in items:
        r: Dict[str, Any] = {}
        for k, v in (o or {}).items():
            fld = _field_of(str(k))
            if not fld:
                continue
            if fld == "seq":
                if _is_seq(v):
                    r["seq"] = int(_norm(v))
            elif fld == "issue_types":
                r["issue_types"] = _split_issues(v)
            else:
                r[fld] = _norm(v)
        if r.get("app_name"):
            r.setdefault("issue_types", [])
            _lift_version(r)
            out.append(r)
    return out


def parse_attachment_ex(path: Path, cache_dir: Optional[Path] = None, llm: Any = None,
                        ocr_cache: Optional[Path] = None,
                        stats: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """分派解析链，返回 {rows, chain, error}（chain 便于诊断"这篇用了哪条链"）。

    纪律：**"解析坏了"与"名单本来就是空的"必须可区分**——每条链失败时都要给出
    可定位的 error（异常类型 / 工具退出码 / 缺失依赖），0 行且 error 非空 = 有问题。
    扩展名统一 `lower()`（政府/Windows CMS 常给 `.PDF`/`.DOCX`）。"""
    path = Path(path)
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            rows = parse_docx(path)
            err = ""
            if not rows:
                try:
                    import docx as _docx
                    d = _docx.Document(str(path))
                    n_tables = len(d.tables)
                    err = (f"docx 打开成功但无表格（tables={n_tables}，可能是纯文本 docx）"
                           if n_tables == 0 else "docx 有表格但没装配出行（版式/列名不在别名表）")
                except Exception as e:
                    err = f"docx 打开失败（文件损坏或不是 docx）: {type(e).__name__}: {e}"
            return {"rows": rows, "chain": "docx", "error": err}
        if suffix == ".doc":
            html, herr = _doc_html_ex(path, cache_dir)
            if html:
                rows = parse_inline_html(html)
                if rows:
                    return {"rows": rows, "chain": "doc", "error": ""}
                herr = "soffice 转出的 HTML 里没有可装配的表格"
            cells, terr = _textutil_cells_ex(path)
            rows = _assemble_doc_rows(cells)
            if rows:
                return {"rows": rows, "chain": "doc-textutil", "error": ""}
            return {"rows": [], "chain": "doc-textutil",
                    "error": f"两条 doc 链都没出结果（soffice: {herr or '无输出'}；"
                             f"textutil: {terr or '有输出但装配不出行'}）"}
        if suffix == ".pdf":
            rows = parse_pdf(path)
            if rows:
                return {"rows": rows, "chain": "pdf", "error": ""}
            pngs, perr = _scanned_pdf_pngs_ex(path, Path(cache_dir or "out/pdf_png"))
            if pngs:
                rows = parse_images(pngs, llm=llm, ocr_cache=ocr_cache, stats=stats)
                return {"rows": rows, "chain": "pdf-scan-ocr",
                        "error": perr or ("" if rows else "扫描件渲染成功但 OCR 没出结果")}
            return {"rows": [], "chain": "pdf",
                    "error": f"PDF 抽不出表格（{perr or '既无文本层也无表格'}）"}
        if suffix in (".png", ".jpg", ".jpeg", ".gif"):
            rows = parse_images([path], llm=llm, ocr_cache=ocr_cache, stats=stats)
            return {"rows": rows, "chain": "img", "error": "" if rows else "图片链没出结果（无行带且整图 VLM 未返回行）"}
        if suffix in (".xlsx", ".xls"):
            # 名单/清单类 Excel：复用 pdf_table 的 xlsx 解析（不重复实现）
            try:
                from . import pdf_table as _pt
                r = _pt.parse_xlsx(str(path))
            except Exception as e:
                return {"rows": [], "chain": "xlsx",
                        "error": f"xlsx 解析异常（.xls/加密文件不受支持）: {type(e).__name__}: {e}"}
            if r.get("kind") == "error":
                return {"rows": [], "chain": "xlsx", "error": str(r.get("error"))[:200]}
            rows = _rows_from_keyed_dicts(r.get("rows") or [])
            return {"rows": rows, "chain": "xlsx",
                    "error": "" if rows else "xlsx 打开成功但没装配出行（列名不在别名表）"}
    except Exception as e:
        return {"rows": [], "chain": "error", "error": f"{type(e).__name__}: {e}"}
    return {"rows": [], "chain": "none", "error": f"不支持的附件类型 {path.suffix}"}


def parse_attachment(path: Path, cache_dir: Optional[Path] = None, llm: Any = None,
                     ocr_cache: Optional[Path] = None) -> List[Dict[str, Any]]:
    """单附件 → 行（不抛异常；要诊断信息用 parse_attachment_ex）。"""
    return parse_attachment_ex(path, cache_dir, llm, ocr_cache)["rows"]


# ---------------- 整篇聚合 ----------------

def _list_signature(rows: Sequence[Dict[str, Any]]) -> set:
    return {_norm(r.get("app_name")) for r in rows if _norm(r.get("app_name"))}


def att_header_text_ex(path: Path, cache_dir: Optional[Path] = None) -> Tuple[str, str]:
    """附件表头文本 + 读取错误。返回 ("", 原因) = 读不到（**不等于**"表头为空"）。"""
    path = Path(path)
    suffix = path.suffix.lower()
    try:
        if suffix == ".doc":
            html, err = _doc_html_ex(path, cache_dir)
            if not html:
                return "", f"soffice 转 HTML 失败: {err}"
            from lxml import html as LH
            d = LH.fromstring(html)
            trs = d.xpath("//table//tr[1]")
            return (" ".join(trs[0].itertext())[:120] if trs else ""), ("" if trs else "HTML 里没有表格")
        if suffix == ".docx":
            import docx
            dd = docx.Document(str(path))
            if not dd.tables:
                return "", "docx 无表格"
            return " ".join(c.text for c in dd.tables[0].rows[0].cells)[:120], ""
        if suffix == ".pdf":
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                t = (pdf.pages[0].extract_tables() or [[]])[0]
                if not t:
                    return "", "PDF 首页无可抽表格"
                return " ".join((c or "") for c in t[0])[:120], ""
        if suffix in (".xlsx", ".xls"):
            from . import pdf_table as _pt
            r = _pt.parse_xlsx(str(path))
            if r.get("kind") == "error":
                return "", str(r.get("error"))[:120]
            rows = r.get("rows") or []
            if rows and isinstance(rows[0], dict):
                return " ".join(str(k) for k in rows[0].keys())[:120], ""
            return "", "xlsx 无表头行"
    except Exception as e:
        return "", f"{type(e).__name__}: {e}"
    return "", f"该类型无表头读取链（{suffix}）"


def att_header_text(path: Path, cache_dir: Optional[Path] = None) -> str:
    """附件表头文本（用于名单分类：识别"复测"类附件）。读不到返回 ""。"""
    return att_header_text_ex(path, cache_dir)[0]


def aggregate_ex(files: Sequence[str], cache_dir: Optional[Path] = None, llm: Any = None,
                 ocr_cache: Optional[Path] = None, dedupe_lists: bool = True,
                 inline_html: Optional[str] = None) -> Dict[str, Any]:
    """`aggregate` 的带诊断版本：返回 {rows, src, files, notes, stats}。

    `files` 逐条记 {name, chain, rows, error}——**0 行与失败也进账**，
    否则"这篇正文本来没有名单"与"5 个附件全挂了"在返回值上同形（违反 0 结果纪律）。
    """
    files_sorted = sorted((Path(x) for x in files), key=lambda p: str(p))
    doc_like = [f for f in files_sorted
                if f.suffix.lower() in (".doc", ".docx", ".pdf", ".xls", ".xlsx", ".wps")]
    img_like = [f for f in files_sorted if f.suffix.lower() in (".png", ".jpg", ".jpeg", ".gif")]
    unknown = [f for f in files_sorted if f not in doc_like and f not in img_like]
    diag: List[Dict[str, Any]] = []
    stats: Dict[str, Any] = {}
    notes: List[str] = []

    def _parse(f: Path) -> List[Dict[str, Any]]:
        d = parse_attachment_ex(f, cache_dir, llm, ocr_cache, stats)
        rows = merge_app_rows(d["rows"])
        diag.append({"name": f.name, "chain": d["chain"], "rows": len(rows),
                     "error": d["error"]})
        return rows

    parsed: List[Tuple[Path, List[Dict[str, Any]]]] = []
    for f in doc_like:
        rows = _parse(f)
        if rows:
            parsed.append((f, rows))
    skipped_images = 0
    if not parsed:
        for f in img_like:
            rows = _parse(f)
            if rows:
                parsed.append((f, rows))
    else:
        skipped_images = len(img_like)
        if skipped_images:
            notes.append(f"文档优先：跳过 {skipped_images} 张图（防与文档名单重复计数）")
    for f in unknown:
        diag.append({"name": f.name, "chain": "none", "rows": 0,
                     "error": f"不在支持表内（{f.suffix or '无扩展名'}）"})
    if not parsed and inline_html:
        rows = merge_app_rows(parse_inline_html(inline_html))
        if rows:
            return {"rows": rows, "src": "inline", "files": diag, "notes": notes, "stats": stats}
    if not parsed:
        # 全失败也要给出可读的 src（否则 wrapper 返回的 (rows, src) 里没有任何缺口痕迹）
        src = " | ".join(f"{d['name']}:0({d['error'] or d['chain']})" for d in diag) if diag else ""
        return {"rows": [], "src": src, "files": diag, "notes": notes, "stats": stats}

    def _artifacts(rows: Sequence[Dict[str, Any]]) -> int:
        return sum(1 for x in rows
                   if _norm(x.get("version")) and _norm(x.get("app_name")).endswith(_norm(x.get("version"))))

    kept: List[Dict[str, Any]] = []
    for f, rows in parsed:
        hdr = att_header_text(f, cache_dir)
        if not hdr:
            notes.append(f"{f.name} 表头读取为空（无法判定是否复测名单）")
        kind = "复测名单（反复出现同类问题企业）" if "复测" in hdr else ""
        if kind:
            for x in rows:
                x["list_kind"] = kind
        for x in rows:
            x.setdefault("source_file", f.name)
        if not dedupe_lists:
            kept.append({"key": _list_signature(rows), "rows": rows, "name": f.name, "art": 0})
            continue
        key = _list_signature(rows)
        hit = None
        for k in kept:
            inter = len(key & k["key"])
            if inter >= 0.9 * max(len(key), len(k["key"]), 1):
                hit = k
                break
        if hit is None:
            kept.append({"key": key, "rows": rows, "name": f.name, "art": _artifacts(rows)})
        elif _artifacts(rows) < hit["art"]:
            notes.append(f"二次上传版本去重：改用 {f.name}（名称粘连更少）")
            hit.update({"rows": rows, "name": f.name, "art": _artifacts(rows)})
        else:
            notes.append(f"二次上传版本去重：跳过 {f.name}（与 {hit['name']} 名称集合重叠≥90%）")
    rows_out: List[Dict[str, Any]] = []
    parts: List[str] = []
    for k in kept:
        rows_out += k["rows"]
        parts.append(f"{k['name']}:{len(k['rows'])}")
    src = "+".join(parts)
    bad = [d for d in diag if d["rows"] == 0]
    if bad:
        src += " | 未出结果: " + ", ".join(f"{d['name']}:0({d['error'] or d['chain']})" for d in bad)
    return {"rows": rows_out, "src": src, "files": diag, "notes": notes, "stats": stats}


def aggregate(files: Sequence[str], cache_dir: Optional[Path] = None, llm: Any = None,
              ocr_cache: Optional[Path] = None, dedupe_lists: bool = True,
              inline_html: Optional[str] = None) -> Tuple[List[Dict[str, Any]], str]:
    """一份通报的全部附件 → 名单行（逐附件解析 + 附件内合并 + 容错去重）。

    **文档优先、图片兜底、内联殿后**：同一篇通报常同时挂 doc 名单与它的截图版
    （总1：doc 41 行 + 镜像补链的几张图），都解析会重复计数（实测 41 → 55）；
    下架/回头看类常把名单直接贴正文（无附件），需调用方传 `inline_html`。

    去重只做**跨附件**的"同一名单二次上传版本"（名称集合重叠 ≥90%），保留
    "名称粘连版本号"更少的一版；附件内不做跨附件合并（同一 APP 出现在不同地区
    名单里是两条独立记录，合并会少计）。返回 (rows, "att_01.doc:39+att_02.doc:46")；
    `src` 里**0 行/失败的附件也会列出**（`| 未出结果: att_03.xlsx:0(原因)`）——
    要结构化诊断用 `aggregate_ex`。

    表头含"复测"的附件是"反复出现同类问题企业"名单（如下架通报的附件7）：其行会打
    `list_kind`，**不计入正文声明数**（声明数只覆盖主名单），记录本身保留。
    """
    d = aggregate_ex(files, cache_dir, llm, ocr_cache, dedupe_lists, inline_html)
    return d["rows"], d["src"]


# ---------------- 声明数对账 ----------------

_CN_DIGITS = {"零": 0, "〇": 0, "一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5,
              "六": 6, "七": 7, "八": 8, "九": 9}
_CN_UNITS = {"十": 10, "百": 100, "千": 1000, "万": 10000}
_CN_CHARS = "".join(_CN_DIGITS) + "".join(_CN_UNITS)


def _cn2int(s: str) -> int:
    """中文数字 → int（支持 零/〇/两 与 十/百/千/万）。

    旧实现按字符累加：`一百零六` → 6（"零"不在字符类里被吞）、`一百四十五` → 1045
    （百位不进位）——声明数错会让整篇对账误判，两者都实测复现过。"""
    total, cur = 0, 0
    for ch in s:
        if ch in _CN_DIGITS:
            cur = _CN_DIGITS[ch]
        elif ch == "万":
            total = (total + cur) * 10000
            cur = 0
        elif ch in _CN_UNITS:
            total += (cur or 1) * _CN_UNITS[ch]
            cur = 0
    return total + cur


def stated_count(text: str, stats: Optional[Dict[str, Any]] = None) -> Optional[int]:
    """正文声明的款数（句子级求和）。

    多附件期正文给多个数（部本级 + 各省管理局分列，如 71+74=145）；同一句里也可能
    并列两个口径（"尚有71款…，各通信管理局检查发现仍有74款…"）——按**句内全部命中**
    求和（旧实现每句只取第一个数，实测少计一半）。

    必须排除：
      - "对…368款APP提出整改要求" —— 累计口径，不是本批名单
      - "对上述共计106款APP进行下架" —— 汇总复述，重复计数（"共计"仅与"上述"同句才排除）
    传 `stats` 时回填 `skipped_values`（越界未计数）与 `skipped_sentences`（命中排除规则）。
    """
    if not text:
        return None
    if "<" in text:                                   # 粗略去标签（正文 HTML）
        text = re.sub(r"<script.*?</script>|<style.*?</style>", " ", text, flags=re.S | re.I)
        text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text)
    total, found = 0, False
    num_pat = re.compile(rf"([0-9]{{1,5}}|[{_CN_CHARS}]{{1,6}})\s*款")
    gate = re.compile(r"(?:尚有|仍有|共发现|检查发现|检测发现|发现|涉及|通报|共计|本批)"
                      r"[^款]{0,30}?$")
    for sent in re.split(r"[。；;]", text):
        if "提出整改要求" in sent or "进行下架" in sent or ("共计" in sent and "上述" in sent):
            if stats is not None:
                stats.setdefault("skipped_sentences", []).append(sent.strip()[:60])
            continue
        for m in num_pat.finditer(sent):
            if not gate.search(sent[:m.start(1)]):
                continue
            raw = m.group(1)
            v = int(raw) if raw.isdigit() else _cn2int(raw)
            if 1 <= v <= 5000:
                total += v
                found = True
            elif stats is not None:
                stats.setdefault("skipped_values", []).append(raw)
    return total if found else None


def declared_fields(header_text: str) -> Optional[set]:
    """附件表头声明了哪些字段（完备率的分母口径：源表没有的列不算缺失）。

    **空表头返回 None**（= "没读到表头"，不是"没有列"）——旧实现返回空集合，
    `audit_completeness` 会把这些记录静默排除出检查分母，报出的完备率虚高。"""
    if not _norm_ws(header_text):
        return None
    txt = _norm_ws(header_text)
    return {fld for fld, words in FIELD_SETS if any(_norm_ws(w) in txt for w in words)}


def audit_counts(pairs: Sequence[Tuple[Any, int, Optional[int]]], tol: float = 0.05) -> List[Tuple[Any, int, Optional[int]]]:
    """行数 vs 声明数：返回超差项 [(批次, 实测, 声明)]（无声明数的跳过）。

    注意：返回空列表有两种含义——"全部达标"或"没有任何可比较项"。验收断言请用
    `audit_counts_ex`（它把 `no_stated` 单列出来），否则"根本没做对账"会假通过。
    """
    bad = []
    for batch, got, stated in pairs:
        if stated and abs(got - stated) > tol * stated:
            bad.append((batch, got, stated))
    return bad


def audit_counts_ex(pairs: Sequence[Tuple[Any, int, Optional[int]]], tol: float = 0.05) -> Dict[str, Any]:
    """结构化对账：{checked, bad, no_stated}。

    `no_stated` = 该篇正文没抽出声明数（关键字门限没命中/计数句被排除）——这是
    **待人工确认**项，不是"通过"。验收写法：
        r = audit_counts_ex(pairs); assert not r["bad"] and not r["no_stated"]
    """
    bad = []
    no_stated = []
    for batch, got, stated in pairs:
        if not stated:
            no_stated.append(batch)
        elif abs(got - stated) > tol * stated:
            bad.append((batch, got, stated))
    return {"checked": len(pairs) - len(no_stated), "bad": bad, "no_stated": no_stated}


def audit_completeness(records: Sequence[Dict[str, Any]], declared_by_source: Dict[Any, Optional[set]]) -> Dict[str, Any]:
    """按源表声明列算非空率：缺失格数 / 检查格数（源表无该列的空值不计）。

    `declared_by_source` 的键是 records 里的来源标识（如 (notice_id, 文件名)）：
    - 键缺失 → 按"五列全声明"从严计（宁可报缺失，不放过）；
    - 值为 `None`（表头读不到）→ 同样从严，并计入 `sources_unreadable`；
    - 值为空集合（表头读到但一列都不认识）→ 也从严（否则该来源静默退出分母）。
    `checked == 0` 时 `rate` 为 None（"一个格都没检查"不是 0% 也不是 100%）。
    """
    fields = ("app_name", "company_name", "version", "store_source", "issue_types")
    checked = 0
    missing: List[Tuple[Any, Any, str, Any]] = []
    unreadable: List[Any] = []
    for r in records:
        key = (r.get("notice_id"), r.get("source_file"))
        if key in declared_by_source and not declared_by_source[key]:
            if key not in unreadable:
                unreadable.append(key)
            declared = set(fields)
        elif key in declared_by_source:
            declared = declared_by_source[key]
        else:
            declared = set(fields)
        for fld in sorted(declared):
            checked += 1
            if not r.get(fld):
                missing.append((r.get("batch_no_total"), r.get("seq"), fld, r.get("source_file")))
    return {"checked": checked, "missing": len(missing),
            "rate": (100.0 - len(missing) * 100.0 / checked) if checked else None,
            "missing_detail": missing[:50], "sources_unreadable": unreadable,
            "reason": "" if checked else ("empty_records" if not records else "no_declared_fields")}


def write_jsonl(rows: Iterable[Dict[str, Any]], path: Path) -> int:
    """按行写 JSONL（UTF-8 不转义中文），返回条数。"""
    n = 0
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
            n += 1
    return n
